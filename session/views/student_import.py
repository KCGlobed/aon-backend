from io import BytesIO

from django.db import transaction
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from rest_framework import status
from rest_framework.views import APIView
from rest_framework.parsers import MultiPartParser, FormParser
from rest_framework.permissions import IsAuthenticated
from rolepermissions.roles import assign_role
from session.emails import queue_session_invites
from session.serializers import *
from users.renderers import UserRenderer
from aon_backend.utils import *
from aon_backend.permissions import RoleOrPermissionCheck

SAMPLE_FILE_NAME = 'student_import_template.xlsx'

SAMPLE_INSTRUCTIONS = [
    'Student Import Template',
    f"Fill in one student per row from row {FIRST_DATA_ROW_NUMBER}. Do not move or rename the column names on row {HEADER_ROW_NUMBER}.",
    f"{', '.join(REQUIRED_COLUMNS)} are required. Email must be unique and Application Id is what the student logs in with. "
    f"Example: Riya | Sharma | riya.sharma@example.com | APP-1001 | 12 MG Road | Pune | Maharashtra | India | 411001",
]

# The template deliberately ships with no data row. An example left sitting in the sheet would be
# imported as a real student by anybody who simply fills in the rows beneath it.


def import_student_users(rows):
    """Writes the parsed roster out. A student already on record has their details refreshed; a new
    one gets a student account that can sit a test straight away."""
    created = []
    updated = []

    for row in rows:
        values = row['values']
        existing = row.get('existing_user')

        if existing is None:
            user = User.objects.create_user(
                email=values['email'],
                first_name=values['first_name'],
                last_name=values['last_name'],
                password=generate_random_password(),
            )
            user.role = User.Student
            # Students sign in with their email, application id and name rather than a password.
            user.email_verified = 1
            user.is_active = True
            for field in ['application_id', 'address', 'city', 'state', 'country', 'pincode']:
                setattr(user, field, values[field])
            user.save()
            assign_role(user, 'Student')
            created.append(user)
            continue

        for field in ['first_name', 'last_name', 'application_id', 'address', 'city', 'state',
                      'country', 'pincode']:
            if values[field]:
                setattr(existing, field, values[field])
        existing.save()
        assign_role(existing, 'Student')
        updated.append(existing)

    return created, updated


class ImportStudentView(APIView):
    """Imports a roster of students from an Excel file.

    Nothing is written unless every row is clean, so a rejected file leaves no half-built students
    behind. Pass 'session' alongside the file to add everybody imported to that session and mail
    them the schedule in the same call.
    """

    renderer_classes = [UserRenderer]
    parser_classes = [MultiPartParser, FormParser]
    permission_classes = [IsAuthenticated,
                          RoleOrPermissionCheck.for_permission_or_roles(
                              "import_student",
                            [SuperAdmin]
                        )]
    def post(self, request, format=None):
        session = None
        session_id = request.data.get('session')
        if session_id:
            session = Session.objects.select_related('assessment').filter(pk=session_id).first()
            if session is None:
                return error_response(message="Session not found!", data={}, status_code=status.HTTP_404_NOT_FOUND)

            if session.state == Session.State.COMPLETED:
                return error_response(message="This session is over and students can no longer be added to it!", data={}, status_code=status.HTTP_400_BAD_REQUEST)

        serializer = StudentImportSerializer(data = request.data)
        serializer.is_valid(raise_exception = True)

        rows = serializer.validated_data['rows']
        errors = serializer.validated_data['errors']

        if errors:
            return error_response(
                message="Import failed. No students were imported. Please correct the errors and upload the file again.",
                data={'total_rows': len(rows) + len(errors), 'error_count': len(errors), 'errors': errors},
                status_code=status.HTTP_400_BAD_REQUEST
            )

        try:
            with transaction.atomic():
                created, updated = import_student_users(rows)
                added = self.add_to_session(session, created + updated) if session else []
        except serializers.ValidationError:
            raise
        except Exception as e:
            logger.error(f"Student import failed: {e}")
            return error_response(
                message="Import failed. No students were imported. Please try again.",
                data={'total_rows': len(rows)},
                status_code=status.HTTP_400_BAD_REQUEST
            )

        data = {
            'total_rows': len(rows),
            'created_count': len(created),
            'updated_count': len(updated),
            'students': StudentImportResultSerializer(created + updated, many=True).data,
        }

        if session:
            data['session'] = SessionDropdownSerializer(session).data
            data['added_to_session'] = len(added)
            data['already_in_session'] = len(created) + len(updated) - len(added)
            data['invite_summary'] = queue_session_invites(added)

        return success_response(
            message=f"{len(created)} student(s) created and {len(updated)} updated successfully!",
            data=data,
            status_code=status.HTTP_201_CREATED
        )

    def add_to_session(self, session, students):
        """Puts the imported students on the session, leaving out anybody already on it so a roster
        can be uploaded again without complaint."""
        fresh = [student for student in students
                 if not session.students.filter(student=student).exists()]
        if not fresh:
            return []

        validate_students_are_free(fresh, session.start_datetime, session.end_datetime,
                                   exclude_session=session)

        SessionStudent.objects.bulk_create([
            SessionStudent(session=session, student=student) for student in fresh
        ])

        return list(session.students.select_related('student').filter(student__in=fresh))


class GetStudentImportSampleView(APIView):
    """The Excel template the import expects, with the column names already on the right row."""

    permission_classes = [IsAuthenticated,
                          RoleOrPermissionCheck.for_permission_or_roles(
                              "import_student",
                            [SuperAdmin]
                        )]
    def get(self, request, format=None):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Students'

        for index, line in enumerate(SAMPLE_INSTRUCTIONS, start=1):
            sheet.cell(row=index, column=1, value=line).font = Font(bold=index == 1, size=13 if index == 1 else 11)

        header_fill = PatternFill('solid', start_color='1F3A8A')
        for index, column in enumerate(EXCEL_COLUMNS, start=1):
            cell = sheet.cell(row=HEADER_ROW_NUMBER, column=index, value=column)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
            sheet.column_dimensions[get_column_letter(index)].width = max(len(column) + 6, 18)

        stream = BytesIO()
        workbook.save(stream)
        workbook.close()

        response = HttpResponse(
            stream.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = f'attachment; filename="{SAMPLE_FILE_NAME}"'
        return response
