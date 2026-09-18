from django.core.exceptions import ValidationError as DjangoValidationError
from django.core.validators import validate_email
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from rest_framework import serializers
from session.models import *
from aon_backend.utils import *


EXCEL_COLUMNS = [
    'First Name',
    'Last Name',
    'Email',
    'Application Id',
    'Address',
    'City',
    'State',
    'Country',
    'Pincode',
]

REQUIRED_COLUMNS = ['First Name', 'Email', 'Application Id']

# The column each value is written to, along with what the User model allows it.
COLUMN_FIELDS = {
    'First Name': ('first_name', 100),
    'Last Name': ('last_name', 100),
    'Email': ('email', 255),
    'Application Id': ('application_id', 100),
    'Address': ('address', 255),
    'City': ('city', 120),
    'State': ('state', 120),
    'Country': ('country', 120),
    'Pincode': ('pincode', 120),
}

ALLOWED_EXTENSIONS = ['.xlsx', '.xlsm']

# Rows 1-3 of the template hold the title and instructions, row 4 holds the column
# names and the students themselves start at row 5.
HEADER_ROW_NUMBER = 4
FIRST_DATA_ROW_NUMBER = HEADER_ROW_NUMBER + 1

MAX_IMPORT_ROWS = 5000


def normalize_header(value):
    return "".join(str(value or "").split()).lower()


def clean_cell(value):
    """Excel hands back numbers, dates and None, so everything becomes a trimmed string. Pincodes
    in particular arrive as floats once Excel has treated them as numbers."""
    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return " ".join(str(value).split())


class StudentImportSerializer(serializers.Serializer):
    """Reads a roster of students out of an Excel file.

    A student already on record is matched by email and has their details refreshed rather than
    being rejected, since the same roster is often uploaded again for a later session.
    """

    file = serializers.FileField(required=True)

    def validate_file(self, value):
        if not value.name.lower().endswith(tuple(ALLOWED_EXTENSIONS)):
            raise serializers.ValidationError(
                "Only Excel files ({}) are supported!".format(", ".join(ALLOWED_EXTENSIONS))
            )

        if value.size == 0:
            raise serializers.ValidationError("The uploaded file is empty!")

        return value

    def validate(self, data):
        try:
            workbook = load_workbook(data['file'], read_only=True, data_only=True)
        except InvalidFileException:
            raise serializers.ValidationError(
                "The uploaded file is not a valid Excel file. Please save it as .xlsx and try again."
            )
        except Exception:
            raise serializers.ValidationError("The uploaded file could not be read. Please check the file and try again.")

        try:
            sheet = workbook.active
            rows = sheet.iter_rows(min_row=HEADER_ROW_NUMBER, values_only=True)

            try:
                header_row = next(rows)
            except StopIteration:
                raise serializers.ValidationError(
                    f"The column names could not be found on row {HEADER_ROW_NUMBER} of the file!"
                )

            header_map = {}
            for index, header in enumerate(header_row):
                key = normalize_header(header)
                if key:
                    header_map.setdefault(key, index)

            missing = [column for column in EXCEL_COLUMNS if normalize_header(column) not in header_map]
            if missing:
                raise serializers.ValidationError(
                    "The following columns are missing from row {} of the file: {}.".format(
                        HEADER_ROW_NUMBER, ", ".join(missing)
                    )
                )

            def cell(row, column):
                index = header_map[normalize_header(column)]
                return clean_cell(row[index]) if index < len(row) else ""

            parsed_rows = []
            errors = []
            seen_emails = {}
            seen_application_ids = {}

            for row_number, row in enumerate(rows, start=FIRST_DATA_ROW_NUMBER):
                if not any(clean_cell(value) for value in row):
                    continue

                if len(parsed_rows) + len(errors) >= MAX_IMPORT_ROWS:
                    raise serializers.ValidationError(
                        f"The file contains more than {MAX_IMPORT_ROWS} rows. Please split it into smaller files."
                    )

                row_errors = []
                values = {}

                for column, (field, max_length) in COLUMN_FIELDS.items():
                    value = cell(row, column)
                    if not value and column in REQUIRED_COLUMNS:
                        row_errors.append(f"'{column}' is required.")
                    elif len(value) > max_length:
                        row_errors.append(f"'{column}' must be at most {max_length} characters.")
                    values[field] = value

                email = values['email'].lower()
                values['email'] = email
                if email:
                    try:
                        validate_email(email)
                    except DjangoValidationError:
                        row_errors.append(f"'Email' is not a valid email address, got '{email}'.")
                    else:
                        if email in seen_emails:
                            row_errors.append(f"'Email' {email} is repeated on row {seen_emails[email]}.")
                        else:
                            seen_emails[email] = row_number

                application_id = values['application_id']
                if application_id:
                    key = application_id.lower()
                    if key in seen_application_ids:
                        row_errors.append(
                            f"'Application Id' {application_id} is repeated on row {seen_application_ids[key]}."
                        )
                    else:
                        seen_application_ids[key] = row_number

                if row_errors:
                    errors.append({'row': row_number, 'email': values['email'], 'errors': row_errors})
                    continue

                parsed_rows.append({'row': row_number, 'values': values})
        finally:
            workbook.close()

        self.check_against_existing_users(parsed_rows, errors)

        if not parsed_rows and not errors:
            raise serializers.ValidationError(
                f"The uploaded file does not contain any students. Students must start from row {FIRST_DATA_ROW_NUMBER}."
            )

        errors.sort(key=lambda error: error['row'])

        data['rows'] = parsed_rows
        data['errors'] = errors

        return data

    def check_against_existing_users(self, parsed_rows, errors):
        """Matches the file against the accounts already on record. An email that belongs to
        somebody who is not a student, and an application id already answering for a different
        person, are both refused - either one would break the student login."""
        emails = [row['values']['email'] for row in parsed_rows]
        application_ids = [row['values']['application_id'] for row in parsed_rows]

        users_by_email = {user.email.lower(): user for user in User.objects.filter(email__in=emails)}
        taken_application_ids = {}
        for user in User.objects.filter(application_id__in=application_ids).exclude(application_id__isnull=True):
            taken_application_ids.setdefault(user.application_id.strip().lower(), user)

        kept_rows = []
        for row in parsed_rows:
            values = row['values']
            row_errors = []

            existing = users_by_email.get(values['email'])
            if existing is not None:
                if existing.role != User.Student:
                    row_errors.append(
                        f"'Email' {values['email']} already belongs to a non-student account and cannot be imported."
                    )
                elif existing.is_deleted:
                    row_errors.append(f"'Email' {values['email']} belongs to a deleted account and cannot be imported.")

            clash = taken_application_ids.get(values['application_id'].lower())
            if clash is not None and (existing is None or clash.pk != existing.pk):
                row_errors.append(
                    f"'Application Id' {values['application_id']} is already in use by {clash.email}."
                )

            if row_errors:
                errors.append({'row': row['row'], 'email': values['email'], 'errors': row_errors})
                continue

            row['existing_user'] = existing
            kept_rows.append(row)

        parsed_rows[:] = kept_rows


class StudentImportResultSerializer(serializers.ModelSerializer):
    full_name = serializers.SerializerMethodField()

    class Meta:
        model = User
        fields = ['id', 'uid', 'first_name', 'last_name', 'full_name', 'email', 'application_id',
                  'address', 'city', 'state', 'country', 'pincode']

    def get_full_name(self, obj):
        return " ".join(filter(None, [obj.first_name, obj.last_name])).strip()
