from decimal import Decimal, InvalidOperation
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from rest_framework import serializers
from content.models import *
from aon_backend.utils import *


EXCEL_COLUMNS = [
    'Section Name',
    'DifficultyLevel',
    'Question Code',
    'Question',
    'Correct Answer',
    'OptionA',
    'OptionB',
    'OptionC',
    'OptionD',
    'OptionE',
    'OptionF',
    'Correct Answer Marks',
    'Negative Marking',
]

OPTION_COLUMNS = ['OptionA', 'OptionB', 'OptionC', 'OptionD', 'OptionE', 'OptionF']

OPTION_LETTERS = {
    'OptionA': 'A',
    'OptionB': 'B',
    'OptionC': 'C',
    'OptionD': 'D',
    'OptionE': 'E',
    'OptionF': 'F',
}

DIFFICULTY_MAP = {
    'easy': Question.DifficultyLevel.EASY,
    '1': Question.DifficultyLevel.EASY,
    'medium': Question.DifficultyLevel.MEDIUM,
    '2': Question.DifficultyLevel.MEDIUM,
    'hard': Question.DifficultyLevel.HARD,
    '3': Question.DifficultyLevel.HARD,
}

ALLOWED_EXTENSIONS = ['.xlsx', '.xlsm']

# Rows 1-3 of the template hold the title and instructions, row 4 holds the column
# names and the questions themselves start at row 5.
HEADER_ROW_NUMBER = 4
FIRST_DATA_ROW_NUMBER = HEADER_ROW_NUMBER + 1

# Values that mean "no negative marking" rather than a number.
NOT_APPLICABLE_VALUES = ['na', 'n/a', 'n.a.', 'nil', 'none', '-']

MAX_IMPORT_ROWS = 5000


def normalize_header(value):
    return "".join(str(value or "").split()).lower()


def clean_cell(value):
    """Excel hands back numbers, dates and None, so everything becomes a trimmed string."""
    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value).strip()


def parse_mark(value, column, row_errors, allow_not_applicable=False):
    value = clean_cell(value)
    if not value:
        return Decimal('0.00')

    if allow_not_applicable and value.lower() in NOT_APPLICABLE_VALUES:
        return Decimal('0.00')

    try:
        mark = Decimal(value)
    except InvalidOperation:
        row_errors.append(f"'{column}' must be a number, got '{value}'.")
        return None

    if mark < 0:
        row_errors.append(f"'{column}' cannot be negative, got '{value}'.")
        return None

    if mark >= Decimal('1000'):
        row_errors.append(f"'{column}' must be less than 1000, got '{value}'.")
        return None

    return mark.quantize(Decimal('0.01'))


class QuestionImportSerializer(serializers.Serializer):
    file = serializers.FileField(required=True)

    def validate_file(self, value):
        if not value.name.lower().endswith(tuple(ALLOWED_EXTENSIONS)):
            raise serializers.ValidationError(
                "Only Excel files ({}) are supported!".format(", ".join(ALLOWED_EXTENSIONS))
            )

        if value.size == 0:
            raise serializers.ValidationError("The uploaded file is empty!")

        return value

    def validate(self, data):
        try:
            workbook = load_workbook(data['file'], read_only=True, data_only=True)
        except InvalidFileException:
            raise serializers.ValidationError(
                "The uploaded file is not a valid Excel file. Please save it as .xlsx and try again."
            )
        except Exception:
            raise serializers.ValidationError("The uploaded file could not be read. Please check the file and try again.")

        try:
            sheet = workbook.active
            rows = sheet.iter_rows(min_row=HEADER_ROW_NUMBER, values_only=True)

            try:
                header_row = next(rows)
            except StopIteration:
                raise serializers.ValidationError(
                    f"The column names could not be found on row {HEADER_ROW_NUMBER} of the file!"
                )

            header_map = {}
            for index, header in enumerate(header_row):
                key = normalize_header(header)
                if key:
                    header_map.setdefault(key, index)

            missing = [column for column in EXCEL_COLUMNS if normalize_header(column) not in header_map]
            if missing:
                raise serializers.ValidationError(
                    "The following columns are missing from row {} of the file: {}.".format(
                        HEADER_ROW_NUMBER, ", ".join(missing)
                    )
                )

            def cell(row, column):
                index = header_map[normalize_header(column)]
                return clean_cell(row[index]) if index < len(row) else ""

            sections_by_name = {section.name.strip().lower(): section for section in Sections.objects.all()}
            existing_codes = {code.lower() for code in Question.objects.values_list('question_code', flat=True)}

            parsed_rows = []
            errors = []
            seen_codes = {}

            for row_number, row in enumerate(rows, start=FIRST_DATA_ROW_NUMBER):
                if not any(clean_cell(value) for value in row):
                    continue

                if len(parsed_rows) + len(errors) >= MAX_IMPORT_ROWS:
                    raise serializers.ValidationError(
                        f"The file contains more than {MAX_IMPORT_ROWS} rows. Please split it into smaller files."
                    )

                row_errors = []

                section_name = cell(row, 'Section Name')
                section = sections_by_name.get(section_name.lower())
                if not section_name:
                    row_errors.append("'Section Name' is required.")
                elif section is None:
                    row_errors.append(f"Section '{section_name}' does not exist.")

                difficulty_value = cell(row, 'DifficultyLevel')
                difficulty_level = DIFFICULTY_MAP.get(difficulty_value.lower())
                if not difficulty_value:
                    row_errors.append("'DifficultyLevel' is required.")
                elif difficulty_level is None:
                    row_errors.append(f"'DifficultyLevel' must be Easy, Medium or Hard, got '{difficulty_value}'.")

                question_code = cell(row, 'Question Code')
                if not question_code:
                    row_errors.append("'Question Code' is required.")
                elif question_code.lower() in existing_codes:
                    row_errors.append(f"Question Code '{question_code}' already exists.")
                elif question_code.lower() in seen_codes:
                    row_errors.append(
                        f"Question Code '{question_code}' is repeated on row {seen_codes[question_code.lower()]}."
                    )
                else:
                    seen_codes[question_code.lower()] = row_number

                question_text = cell(row, 'Question')
                if not question_text:
                    row_errors.append("'Question' is required.")

                options = []
                for column in OPTION_COLUMNS:
                    option_text = cell(row, column)
                    if option_text:
                        options.append({'letter': OPTION_LETTERS[column], 'option_text': option_text})

                if len(options) < 2:
                    row_errors.append("At least 2 options are required.")

                correct_answer = cell(row, 'Correct Answer')
                correct_letter = None
                if not correct_answer:
                    row_errors.append("'Correct Answer' is required.")
                else:
                    letters = {option['letter']: option for option in options}
                    if correct_answer.upper() in OPTION_LETTERS.values():
                        if correct_answer.upper() in letters:
                            correct_letter = correct_answer.upper()
                        else:
                            row_errors.append(
                                f"'Correct Answer' is '{correct_answer}' but Option{correct_answer.upper()} is empty."
                            )
                    else:
                        matches = [option for option in options
                                   if option['option_text'].lower() == correct_answer.lower()]
                        if len(matches) == 1:
                            correct_letter = matches[0]['letter']
                        elif not matches:
                            row_errors.append(
                                f"'Correct Answer' value '{correct_answer}' does not match any option. "
                                f"Use a letter A-F or the exact option text."
                            )
                        else:
                            row_errors.append(
                                f"'Correct Answer' value '{correct_answer}' matches more than one option."
                            )

                correct_answer_mark = parse_mark(cell(row, 'Correct Answer Marks'), 'Correct Answer Marks', row_errors)
                negative_mark = parse_mark(cell(row, 'Negative Marking'), 'Negative Marking', row_errors,
                                           allow_not_applicable=True)

                if row_errors:
                    errors.append({'row': row_number, 'question_code': question_code, 'errors': row_errors})
                    continue

                parsed_rows.append({
                    'section': section,
                    'question_code': question_code,
                    'question_text': question_text,
                    'difficulty_level': difficulty_level,
                    'correct_answer_mark': correct_answer_mark,
                    'negative_mark': negative_mark,
                    'options': [
                        {'option_text': option['option_text'], 'right_option': option['letter'] == correct_letter}
                        for option in options
                    ],
                })
        finally:
            workbook.close()

        if not parsed_rows and not errors:
            raise serializers.ValidationError(
                f"The uploaded file does not contain any questions. Questions must start from row {FIRST_DATA_ROW_NUMBER}."
            )

        data['rows'] = parsed_rows
        data['errors'] = errors

        return data
